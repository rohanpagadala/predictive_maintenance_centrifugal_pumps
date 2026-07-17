from __future__ import annotations

import html
import json
import os
import sys
from io import BytesIO, StringIO
from pathlib import Path

import pandas as pd
import plotly.express as px
import requests
import streamlit as st
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import model_persistence as mp
import model_registry as mr
import pdm_utils as u
import retrain
from app import config

API_URL = os.getenv("PDM_API_URL", "http://localhost:8000")

ACCENT = "#0F6E77"
GOOD, GOOD_BG = "#1E8E3E", "#E7F6EC"
WARN, WARN_BG = "#E8A33D", "#FCEFDA"
CRIT, CRIT_BG = "#C23636", "#FBEAEA"
INK_ON_WARN = "#1B2430"

FAULT_COLORS = {
    "Bearing_Failure": "#7C5CBF", "Motor_Failure": "#C2410C",
    "Seal_Failure": "#2C7DA0", "Unknown_Critical_Fault": "#6B7280",
}
CLASS_COLORS = {"Normal": GOOD, "Warning": WARN, "Critical": CRIT, **FAULT_COLORS}

CRITICAL_RUL_THRESHOLD_HOURS = 24
WARNING_RUL_THRESHOLD_HOURS = 360
WARNING_RUL_THRESHOLD_DAYS = WARNING_RUL_THRESHOLD_HOURS // 24

MAX_ONSCREEN_TABLE_ROWS = 500

st.set_page_config(page_title="PdM Dashboard", page_icon="🛠️", layout="wide")


def inject_custom_css() -> None:
    st.markdown(
        f"""
        <style>
          html, body, [class*="st-"] {{ font-variant-numeric: tabular-nums; }}

          .pdm-header {{ padding-bottom: 6px; margin-bottom: 18px; border-bottom: 1px solid rgba(128,128,128,.25); }}
          .pdm-header .pdm-title {{ font-size: 1.9rem; font-weight: 800; letter-spacing: -.01em; margin: 0; }}
          .pdm-header .pdm-caption {{ color: rgba(128,128,128,.95); font-size: .92rem; margin-top: 4px; }}

          .pdm-stat-row {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(170px, 1fr)); gap: 12px; margin: 4px 0 18px; }}
          .pdm-stat {{
            border: 1px solid rgba(128,128,128,.25); border-radius: 10px; padding: 14px 16px;
            background: rgba(128,128,128,.05);
          }}
          .pdm-stat.accent {{ border-color: rgba(15,110,119,.45); background: rgba(15,110,119,.07); }}
          .pdm-stat.good {{ border-color: rgba(30,142,62,.4); background: rgba(30,142,62,.08); }}
          .pdm-stat.warn {{ border-color: rgba(232,163,61,.5); background: rgba(232,163,61,.10); }}
          .pdm-stat.crit {{ border-color: rgba(194,54,54,.45); background: rgba(194,54,54,.08); }}
          .pdm-stat .pdm-stat-value {{ font-size: 1.55rem; font-weight: 800; line-height: 1.15; }}
          .pdm-stat .pdm-stat-label {{ font-size: .74rem; text-transform: uppercase; letter-spacing: .04em; color: rgba(128,128,128,.95); margin-top: 3px; }}

          .pdm-sidebar-status {{
            display: flex; align-items: center; gap: 8px; font-size: .82rem;
            padding: 8px 10px; border-radius: 8px; background: rgba(128,128,128,.08);
            border: 1px solid rgba(128,128,128,.2); margin-bottom: 10px;
          }}
          .pdm-dot {{ width: 8px; height: 8px; border-radius: 50%; flex: none; }}
          .pdm-dot.online {{ background: {GOOD}; box-shadow: 0 0 0 3px rgba(30,142,62,.18); }}
          .pdm-dot.offline {{ background: {CRIT}; box-shadow: 0 0 0 3px rgba(194,54,54,.18); }}

          .pdm-badge {{ display:inline-block; padding: 3px 12px; border-radius: 999px; font-weight: 700; font-size: .85rem; color: #fff; }}

          [data-testid="stDataFrame"], [data-testid="stTable"] {{ border-radius: 8px; overflow: hidden; }}
        </style>
        """,
        unsafe_allow_html=True,
    )


def page_header(title: str, caption: str = "", icon: str = "") -> None:
    heading = f"{icon} {title}".strip()
    caption_html = f'<div class="pdm-caption">{html.escape(caption)}</div>' if caption else ""
    st.markdown(
        f'<div class="pdm-header"><p class="pdm-title">{heading}</p>{caption_html}</div>',
        unsafe_allow_html=True,
    )


def stat_card_row(cards: list[dict]) -> None:
    tiles = "".join(
        f'<div class="pdm-stat {c.get("tone", "")}">'
        f'<div class="pdm-stat-value">{html.escape(str(c["value"]))}</div>'
        f'<div class="pdm-stat-label">{html.escape(str(c["label"]))}</div>'
        f"</div>"
        for c in cards
    )
    st.markdown(f'<div class="pdm-stat-row">{tiles}</div>', unsafe_allow_html=True)


@st.cache_resource(show_spinner="Loading model artifacts...")
def load_artifacts() -> mp.ModelArtifacts:
    return mr.load_models_from_registry(config.MODELS_DIR)


@st.cache_data(ttl=5, show_spinner=False)
def api_health() -> dict | None:
    try:
        resp = requests.get(f"{API_URL}/health", timeout=3)
        resp.raise_for_status()
        return resp.json()
    except requests.RequestException:
        return None


@st.cache_data(ttl=30, show_spinner=False)
def get_active_model_version() -> str:
    try:
        resp = requests.get(f"{API_URL}/models", timeout=5)
        resp.raise_for_status()
        for version in resp.json().get("versions", []):
            if version.get("is_stable"):
                return version["version"]
    except (requests.RequestException, KeyError, ValueError):
        pass
    return "unknown"


def format_api_error(detail) -> str:
    if isinstance(detail, str):
        return detail
    if not isinstance(detail, list):
        return str(detail)

    lines = []
    for err in detail:
        if not isinstance(err, dict) or "msg" not in err:
            lines.append(str(err))
            continue
        loc = err.get("loc", [])
        field = next((str(p) for p in reversed(loc) if isinstance(p, str)), "input")
        line = f"**{field}**: {err['msg']}"
        if "input" in err and err["input"] is not None:
            line += f" (you entered `{err['input']}`)"
        lines.append(line)
    return "\n\n".join(f"- {line}" for line in lines)


def call_predict(payload: dict) -> dict:
    resp = requests.post(f"{API_URL}/predict", json=payload, timeout=15)
    if resp.status_code != 200:
        raise RuntimeError(format_api_error(resp.json().get("detail", resp.text)))
    return resp.json()


def call_explain(payload: dict) -> dict:
    resp = requests.post(f"{API_URL}/explain", json=payload, timeout=30)
    if resp.status_code != 200:
        raise RuntimeError(format_api_error(resp.json().get("detail", resp.text)))
    return resp.json()


def call_predict_csv(file_bytes: bytes, filename: str) -> pd.DataFrame:
    resp = requests.post(
        f"{API_URL}/predict-csv",
        files={"file": (filename, file_bytes, "text/csv")},
        timeout=60,
    )
    if resp.status_code != 200:
        detail = resp.json().get("detail", resp.text) if resp.headers.get("content-type") == "application/json" else resp.text
        raise RuntimeError(format_api_error(detail))
    return pd.read_csv(StringIO(resp.text))


def state_badge(state: str) -> str:
    color = CLASS_COLORS.get(state, "#6B7280")
    text_color = INK_ON_WARN if color == WARN else "#fff"
    label = html.escape(str(state).replace("_", " "))
    return f"<span class='pdm-badge' style='background:{color};color:{text_color}'>{label}</span>"


def format_timestamp(value) -> str:
    try:
        ts = pd.to_datetime(value)
    except (ValueError, TypeError):
        return str(value)
    if ts.tzinfo is None:
        ts = ts.tz_localize("UTC")
    ts = ts.tz_convert("Asia/Kolkata")
    return ts.strftime("%Y-%m-%d %H:%M:%S IST")


def log_history_entry(entry: dict) -> None:
    if "history" not in st.session_state:
        st.session_state["history"] = []
    st.session_state["history"].insert(0, entry)


def estimated_replacement_date(rul_hours) -> str:
    if pd.isnull(rul_hours):
        return ""
    return format_timestamp(pd.Timestamp.now(tz="UTC") + pd.Timedelta(hours=float(rul_hours)))


def rul_severity(rul) -> str:
    if pd.isnull(rul):
        return "unknown"
    if rul <= CRITICAL_RUL_THRESHOLD_HOURS:
        return "critical"
    if rul <= WARNING_RUL_THRESHOLD_HOURS:
        return "warning"
    return "normal"


def recommended_action(rul) -> str:
    severity = rul_severity(rul)
    if severity == "critical":
        return "Immediate Replacement Required"
    if severity == "warning":
        return f"Schedule Replacement Within {WARNING_RUL_THRESHOLD_DAYS} Days"
    return ""


def _rul_row_style(row: pd.Series) -> list[str]:
    severity = rul_severity(row.get("RUL_Hours_Predicted"))
    if severity == "critical":
        return [f"background-color: {CRIT}; color: white"] * len(row)
    if severity == "warning":
        return [f"background-color: {WARN}; color: {INK_ON_WARN}"] * len(row)
    return [""] * len(row)


def priority_icon(rul) -> str:
    severity = rul_severity(rul)
    if severity == "critical":
        return "🔴 Immediate"
    if severity == "warning":
        return f"🟡 {WARNING_RUL_THRESHOLD_DAYS}-day"
    if severity == "normal":
        return "🟢 OK"
    return ""


def add_priority_column(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.insert(1 if "Asset_ID" in df.columns else 0, "Priority", df["RUL_Hours_Predicted"].apply(priority_icon))
    return df


LIKELY_FAILURE_INFO = {
    "Bearing_Failure": (
        "Bearing Failure", "High vibration and temperature indicate possible bearing wear.",
        "Schedule bearing inspection and replacement before failure.",
    ),
    "Seal_Failure": (
        "Seal Failure", "Pressure loss and reduced flow suggest a possible seal leak.",
        "Inspect and replace the mechanical seal; check for leakage.",
    ),
    "Motor_Failure": (
        "Motor Failure", "Elevated temperature and reduced performance indicate a possible motor issue.",
        "Inspect motor windings and cooling; schedule motor service.",
    ),
    "Unknown_Critical_Fault": (
        "Unknown Critical Fault",
        "Sensor readings are critical but don't clearly match a single known failure pattern -- further inspection recommended.",
        "Perform a manual inspection to identify the root cause.",
    ),
    "Cavitation": ("Cavitation", "Pressure fluctuations and vibration suggest cavitation.",
                    "Check suction conditions (NPSH) and reduce inlet throttling."),
    "Impeller_Damage": ("Impeller Damage", "Reduced flow and pressure indicate possible impeller wear.",
                         "Inspect the impeller for wear or damage; schedule replacement."),
    "Shaft_Misalignment": ("Shaft Misalignment", "Excessive vibration suggests shaft alignment issues.",
                            "Realign the shaft and coupling; verify alignment tolerances."),
    "Rotor_Imbalance": ("Rotor Imbalance", "Persistent vibration indicates possible rotor imbalance.",
                         "Perform rotor balancing; inspect for debris or wear causing imbalance."),
    "Lubrication_Failure": ("Lubrication Failure", "Increasing vibration and temperature suggest insufficient lubrication.",
                             "Check lubricant level and quality; replenish or replace."),
    "Coupling_Failure": ("Coupling Failure", "Power transfer irregularities indicate a possible coupling issue.",
                          "Inspect the coupling for wear or damage; replace as needed."),
    "Suction_Blockage": ("Suction Blockage", "Low inlet pressure and reduced flow suggest a suction-side blockage.",
                          "Clear the suction line obstruction; inspect the strainer/filter."),
    "Discharge_Blockage": ("Discharge Blockage", "Rising discharge pressure with falling flow suggests a discharge-side blockage.",
                            "Clear the discharge line obstruction; inspect the discharge valve."),
    "Valve_Malfunction": ("Valve Malfunction", "Unstable pressure and flow suggest a valve malfunction.",
                           "Inspect and service the control/discharge valve."),
    "Overheating": ("Overheating", "Extremely high temperature indicates the pump is overheating.",
                     "Investigate the cooling system and load; allow the pump to cool before restart."),
    "Corrosion_Wear": ("Corrosion / Wear", "Gradual pressure and flow decline suggest corrosion or general wear.",
                        "Inspect wetted components for corrosion; plan refurbishment."),
}

DEVELOPING_ISSUE_INFO = {
    "Bearing_Failure": ("Bearing Wear", "Elevated vibration suggests early bearing wear. Monitor the pump and schedule inspection."),
    "Seal_Failure": ("Seal Degradation", "Early pressure and flow deviation suggests developing seal degradation. Monitor the pump and schedule inspection."),
    "Motor_Failure": ("Motor Stress", "Rising temperature suggests developing motor stress. Monitor the pump and schedule inspection."),
}
UNDETERMINED_DEVELOPING_ISSUE_TEXT = "Unable to determine a dominant degradation pattern."


UNDETERMINED_LIKELY_LABEL = "Not clearly determined"


def _diagnosis_parts(health_state: str, fault_value) -> tuple[str, str, str, str, str] | None:
    if health_state == "Normal":
        return "Likely Condition", GOOD, "Normal Operation", "No abnormal operating conditions detected.", ""
    if health_state == "Warning":
        key = str(fault_value).strip() if pd.notna(fault_value) and str(fault_value).strip() else ""
        entry = DEVELOPING_ISSUE_INFO.get(key)
        if entry is None:
            return "Likely Developing Issue", WARN, UNDETERMINED_LIKELY_LABEL, UNDETERMINED_DEVELOPING_ISSUE_TEXT, ""
        name, sentence = entry
        return "Likely Developing Issue", WARN, name, sentence, ""
    if health_state == "Critical":
        key = str(fault_value).strip() if pd.notna(fault_value) and str(fault_value).strip() else ""
        name, description, action = LIKELY_FAILURE_INFO.get(key, (key.replace("_", " ") or "Unknown Critical Fault", "", ""))
        return "Likely Failure", CRIT, name, description, action
    return None


def _diagnosis_cell_html(health_state: str, fault_value, *, compact: bool = False) -> str:
    parsed = _diagnosis_parts(health_state, fault_value)
    if parsed is None:
        return "—"
    label, color, name, description, action = parsed
    lines = [line for line in (description, f"Recommended action: {action}" if action else "") if line]

    label_style = "font-size:0.72em;text-transform:uppercase;letter-spacing:.04em;opacity:.7"
    parts = [f"<span style='{label_style}'>{html.escape(label)}</span>"]
    if name:
        parts.append(f"<span style='color:{color};font-weight:700'>{html.escape(name)}</span>")
    for line in lines:
        parts.append(f"<span style='font-size:0.85em;opacity:.85'>{html.escape(line)}</span>")
    return "<br>".join(parts)


def add_diagnosis_column(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    health = df.get("Failure_State_Predicted")
    diagnosed = df.get("Diagnosed_Fault")
    if health is None:
        df["Diagnosis"] = "—"
        return df

    fault_col = diagnosed if diagnosed is not None else pd.Series([None] * len(df), index=df.index)
    df["Diagnosis"] = [
        _diagnosis_cell_html(state, fault) for state, fault in zip(health.astype(str), fault_col)
    ]
    return df


def add_diagnosis_export_columns(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    health = df.get("Failure_State_Predicted")
    diagnosed = df.get("Diagnosed_Fault")
    if health is None:
        df["Likely_Failure_Or_Developing_Issue"] = "—"
        df["Failure_Description"] = "—"
        df["Recommended_Action"] = "—"
        return df

    fault_col = diagnosed if diagnosed is not None else pd.Series([None] * len(df), index=df.index)
    names, descriptions, actions = [], [], []
    for state, fault in zip(health.astype(str), fault_col):
        parsed = _diagnosis_parts(state, fault)
        if parsed is None:
            names.append("—")
            descriptions.append("—")
            actions.append("—")
        else:
            _label, _color, name, description, action = parsed
            names.append(name or UNDETERMINED_LIKELY_LABEL)
            descriptions.append(description)
            actions.append(action)
    df["Likely_Failure_Or_Developing_Issue"] = names
    df["Failure_Description"] = descriptions
    df["Recommended_Action"] = actions
    return df


def _format_cell(value) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, float):
        return f"{value:.2f}"
    return str(value)


def cap_for_onscreen_display(display_df: pd.DataFrame) -> tuple[pd.DataFrame, str]:
    if len(display_df) <= MAX_ONSCREEN_TABLE_ROWS or "RUL_Hours_Predicted" not in display_df.columns:
        return display_df, ""
    notice = (
        f"Showing the {MAX_ONSCREEN_TABLE_ROWS:,} most urgent of {len(display_df):,} rows "
        "(lowest RUL first) so the table stays responsive -- every row is still included in "
        "the downloads below."
    )
    capped = display_df.sort_values("RUL_Hours_Predicted", ascending=True).head(MAX_ONSCREEN_TABLE_ROWS)
    return capped, notice


def render_results_table_with_diagnosis(df: pd.DataFrame, height_px: int = 350) -> None:
    header_html = "".join(f"<th>{html.escape(str(col))}</th>" for col in df.columns)

    row_chunks = []
    for i, (_, row) in enumerate(df.iterrows()):
        row_style = "; ".join(_rul_row_style(row)[:1]).strip()
        if row_style:
            tr_style = f" style='{row_style}'"
        else:
            tr_style = " class='stripe'" if i % 2 == 1 else ""
        cells = []
        for col in df.columns:
            value = row[col]
            if col == "Diagnosis":
                cells.append(f"<td class='wrap'>{value}</td>")
            else:
                text = _format_cell(value)
                cells.append(f"<td title='{html.escape(text)}'>{html.escape(text)}</td>")
        row_chunks.append(f"<tr{tr_style}>{''.join(cells)}</tr>")

    table_html = f"""
    <div style="max-height:{height_px}px; overflow:auto; border:1px solid rgba(128,128,128,.3); border-radius:8px;">
      <table class="pdm-results-table" style="width:100%; border-collapse:collapse; font-size:.88rem;">
        <thead>
          <tr>{header_html}</tr>
        </thead>
        <tbody>
          {''.join(row_chunks)}
        </tbody>
      </table>
    </div>
    <style>
      .pdm-results-table td, .pdm-results-table th {{
        padding: 8px 14px; border-bottom: 1px solid rgba(128,128,128,.2); text-align: left;
        white-space: nowrap; max-width: 240px; overflow: hidden; text-overflow: ellipsis;
      }}
      .pdm-results-table td.wrap {{ white-space: normal; max-width: 340px; overflow: visible; text-overflow: clip; }}
      .pdm-results-table thead th {{
        position: sticky; top: 0; background: rgba(128,128,128,.16); backdrop-filter: blur(2px);
        text-transform: uppercase; font-size: .72rem; letter-spacing: .04em; font-weight: 700;
      }}
      .pdm-results-table tbody tr.stripe {{ background: rgba(128,128,128,.06); }}
      .pdm-results-table tbody tr:hover {{ background: rgba(15,110,119,.10); }}
    </style>
    """
    st.markdown(table_html, unsafe_allow_html=True)


def build_replacement_schedule(df: pd.DataFrame) -> pd.DataFrame:
    scheduled = df[df["RUL_Hours_Predicted"] <= WARNING_RUL_THRESHOLD_HOURS].copy()
    scheduled = scheduled.sort_values("RUL_Hours_Predicted", ascending=True)
    scheduled["Recommended_Action"] = scheduled["RUL_Hours_Predicted"].apply(recommended_action)
    scheduled["Estimated_Replacement_Date"] = scheduled["RUL_Hours_Predicted"].apply(estimated_replacement_date)

    column_map = {
        "Asset_ID": "Asset ID",
        "Machine_Model": "Machine Model",
        "Location": "Location",
        "Failure_State_Predicted": "Health Status",
        "Confidence": "Confidence",
        "RUL_Hours_Predicted": "Predicted RUL (Hours)",
        "Estimated_Replacement_Date": "Estimated Replacement Date",
        "Diagnosed_Fault": "Diagnosed Fault",
        "Recommended_Action": "Recommended Action",
    }
    available_cols = [col for col in column_map if col in scheduled.columns]
    return scheduled[available_cols].rename(columns=column_map)


def build_predictions_excel(df: pd.DataFrame) -> bytes:
    df = df.copy()
    for column in df.columns:
        if pd.api.types.is_datetime64_any_dtype(df[column]) and df[column].dt.tz is not None:
            df[column] = df[column].dt.tz_localize(None)

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="predictions")
        worksheet = writer.sheets["predictions"]

        red_fill = PatternFill(start_color=CRIT.lstrip("#"), end_color=CRIT.lstrip("#"), fill_type="solid")
        white_font = Font(color="FFFFFF")
        yellow_fill = PatternFill(start_color=WARN.lstrip("#"), end_color=WARN.lstrip("#"), fill_type="solid")
        black_font = Font(color=INK_ON_WARN.lstrip("#"))
        header_fill = PatternFill(start_color="D7E9EA", end_color="D7E9EA", fill_type="solid")
        header_font = Font(bold=True, color=ACCENT.lstrip("#"))

        n_cols = len(df.columns)
        for col_idx in range(1, n_cols + 1):
            header_cell = worksheet.cell(row=1, column=col_idx)
            header_cell.fill = header_fill
            header_cell.font = header_font
        rul_values = df["RUL_Hours_Predicted"]
        for row_offset, rul in enumerate(rul_values):
            severity = rul_severity(rul)
            if severity == "critical":
                fill, font = red_fill, white_font
            elif severity == "warning":
                fill, font = yellow_fill, black_font
            else:
                continue
            excel_row = row_offset + 2
            for col_idx in range(1, n_cols + 1):
                cell = worksheet.cell(row=excel_row, column=col_idx)
                cell.fill = fill
                cell.font = font

        worksheet.freeze_panes = "A2"
        for col_idx, column in enumerate(df.columns, start=1):
            content_len = int(df[column].astype(str).str.len().fillna(0).max()) if len(df) else 0
            max_len = max(len(str(column)), content_len)
            worksheet.column_dimensions[worksheet.cell(row=1, column=col_idx).column_letter].width = min(max_len + 2, 40)
    return buffer.getvalue()


def page_home():
    page_header(
        "Predictive Maintenance Dashboard",
        caption="Industrial centrifugal pump failure classification & remaining-useful-life prediction.",
    )

    health = api_health()
    if health:
        stat_card_row([
            {"label": "API status", "value": "🟢 Online", "tone": "good"},
            {"label": "Classifier", "value": health.get("classifier_name", "—"), "tone": "accent"},
            {"label": "Regressor", "value": health.get("regressor_name", "—"), "tone": "accent"},
            {"label": "Uptime (s)", "value": f"{health.get('uptime_seconds', 0):.0f}", "tone": "accent"},
        ])
    else:
        stat_card_row([{"label": "API status", "value": "🔴 Offline", "tone": "crit"}])
        st.warning(f"Could not reach API at `{API_URL}`. Start it with:\n\n`uvicorn app.api:app --reload`")

    st.divider()
    st.subheader("What this dashboard does")
    nav_col1, nav_col2, nav_col3 = st.columns(3)
    with nav_col1.container(border=True):
        st.markdown("**Single Prediction**")
        st.caption("Score one pump reading (optionally with SHAP explanation) and log it to history.")
    with nav_col2.container(border=True):
        st.markdown("**Batch Prediction**")
        st.caption("Upload a CSV or Excel file of pump readings and download predictions for every row.")
    with nav_col3.container(border=True):
        st.markdown("**Prediction History**")
        st.caption("A running log of this session's predictions, plus the full persisted audit trail.")


def page_overview():
    page_header("Project Overview", caption="What this system does, and where its limits are.")

    st.markdown(
        """
        ### Business problem
        Unplanned failures of industrial centrifugal pumps cause costly downtime.
        This project turns hourly pump telemetry (vibration, temperature, pressure,
        flow rate) into two decision-support signals:

        1. **Failure classification** — Normal / Warning / Critical
        2. **Remaining Useful Life (RUL) regression** — hours until failure.

        ### Pipeline stages
        Dataset Understanding → EDA → Preprocessing → Feature Engineering →
        Feature Selection → Data Preparation → Model Training (Decision Tree,
        Random Forest, XGBoost, LightGBM) → Evaluation → Comparison → Persistence →
        **this serving layer**.

        ### Architecture: health staging vs. fault diagnosis
        The classifier only stages machine health (Normal/Warning/Critical) —
        it no longer predicts *which* component is failing. When a reading is
        Warning or Critical, a separate rule-based module (`app/fault_diagnosis.py`)
        estimates whether the sensor pattern most resembles a Bearing, Motor,
        or Seal issue, or reports "Unknown_Critical_Fault" when no single
        channel's deviation clearly dominates -- framed as a developing issue to
        watch at Warning severity, or a likely failure at Critical severity. This
        keeps health staging (which the data supports well) separate from fault
        typing (which it doesn't, yet) — see **About** for why.
        """
    )
    st.info(
        "**Known limitation, stated plainly** — the historical Bearing/Motor/Seal_Failure examples "
        "total only ~20 rows in the entire training set, concentrated in a single pump, with "
        "near-identical sensor signatures across the three failure modes. That's why fault typing is "
        "a documented heuristic rule engine rather than a trained classifier — see **About** for the "
        "full explanation."
    )

    st.subheader("Selected feature set")
    try:
        artifacts = load_artifacts()
        with st.container(border=True):
            st.caption(f"{len(artifacts.feature_list)} engineered features feed both models, in this fixed order.")
            st.dataframe(pd.DataFrame({"Feature": artifacts.feature_list}), width="stretch", height=300, hide_index=True)
    except Exception as exc:
        st.error(f"Could not load model artifacts: {exc}")


def get_known_machine_models() -> list[str]:
    fallback = ["Nova-P", "Titan-X3", "Titan-X2", "Atlas-7", "Hydra-500"]
    try:
        artifacts = load_artifacts()
    except Exception as exc:
        st.warning(f"Could not load model artifacts ({exc}) -- showing a fallback machine model list.")
        return fallback
    encoder = artifacts.encoders.encoders.get("Machine_Model")
    if encoder is None:
        st.warning("Model artifacts have no 'Machine_Model' encoder -- showing a fallback list.")
        return fallback
    return sorted(encoder.classes_.tolist())


def get_known_locations() -> list[str]:
    fallback = ["CPF-1", "Offshore-Platform-A", "Offshore-Platform-B", "Onshore-Terminal-1", "Onshore-Terminal-2"]
    try:
        load_artifacts()
    except Exception as exc:
        st.warning(f"Could not load model artifacts ({exc}) -- showing a fallback location list.")
    return fallback


def page_single_prediction():
    page_header(
        "Single Prediction",
        caption="Enter one pump reading -- the result below updates automatically as you change values. "
                "Click Predict to save that result to Prediction History.",
    )

    with st.container(border=True):
        st.markdown("**Pump reading**")
        c1, c2, c3 = st.columns(3)
        asset_id = c1.text_input("Asset ID", value="PUMP_DEMO", help="Free-text identifier for this pump, e.g. its tag number.")
        machine_model = c2.selectbox("Machine Model", options=get_known_machine_models(), help="Pump model, used as a categorical feature by both models.")
        location = c3.selectbox("Location", options=get_known_locations(), help="Site or facility this pump is installed at.")

        vib_lo, vib_hi = u.VALID_RANGES["Vibration_mm_s"]
        temp_lo, temp_hi = u.VALID_RANGES["Temperature_C"]
        pres_lo, pres_hi = u.VALID_RANGES["Pressure_psi"]
        flow_lo, flow_hi = u.VALID_RANGES["Flow_Rate_m3_h"]

        c1, c2, c3, c4 = st.columns(4)
        vibration = c1.number_input(
            f"Vibration (mm/s) [{vib_lo}-{vib_hi}]", min_value=float(vib_lo), max_value=float(vib_hi),
            value=2.2, step=0.1, help="Casing vibration velocity. Rising vibration is the leading signal for bearing wear and misalignment.",
        )
        temperature = c2.number_input(
            f"Temperature (°C) [{temp_lo}-{temp_hi}]", min_value=float(temp_lo), max_value=float(temp_hi),
            value=62.0, step=0.5, help="Bearing/casing temperature. Sustained highs point toward motor or lubrication issues.",
        )
        pressure = c3.number_input(
            f"Pressure (psi) [{pres_lo}-{pres_hi}]", min_value=float(pres_lo), max_value=float(pres_hi),
            value=145.0, step=1.0, help="Discharge pressure. A drop often accompanies seal leakage.",
        )
        flow_rate = c4.number_input(
            f"Flow Rate (m³/h) [{flow_lo}-{flow_hi}]", min_value=float(flow_lo), max_value=float(flow_hi),
            value=250.0, step=1.0, help="Volumetric flow rate. Falling flow alongside falling pressure reinforces a seal-failure pattern.",
        )

        explain = st.checkbox(
            "Explain this prediction (SHAP)", value=False,
            help="Slower than a bare prediction -- computes SHAP values and waterfall plots for both models. "
                 "Recomputes automatically on every value change while checked.",
        )

    current_inputs = (asset_id, machine_model, location, vibration, temperature, pressure, flow_rate, explain)
    if st.session_state.get("single_prediction_inputs") != current_inputs:
        payload = {
            "readings": [{
                "asset_id": asset_id,
                "machine_model": machine_model,
                "location": location,
                "vibration_mm_s": vibration,
                "temperature_c": temperature,
                "pressure_psi": pressure,
                "flow_rate_m3_h": flow_rate,
            }]
        }
        try:
            with st.spinner("Computing SHAP explanation..." if explain else "Scoring reading..."):
                if explain:
                    explanation = call_explain(payload)
                    result = explanation["prediction"]
                else:
                    explanation = None
                    result = call_predict(payload)
        except Exception as exc:
            st.error(f"Prediction failed: {exc}")
            return
        st.session_state["single_prediction_inputs"] = current_inputs
        st.session_state["single_prediction_result"] = result
        st.session_state["single_prediction_explanation"] = explanation
        st.session_state["single_prediction_saved"] = False

    result = st.session_state.get("single_prediction_result")
    if result is None:
        return
    explanation = st.session_state.get("single_prediction_explanation")

    st.divider()
    state = result["failure_state"]
    tone = {"Normal": "good", "Warning": "warn", "Critical": "crit"}.get(state, "accent")
    r1, r2 = st.columns([2, 1])
    with r1:
        st.markdown(f"**Health status**<br>{state_badge(state)}", unsafe_allow_html=True)
        action = recommended_action(result["remaining_useful_life"])
        if action:
            (st.error if tone == "crit" else st.warning)(f"**Recommended action:** {action}")
    with r2:
        if st.button("Predict", type="primary", width="stretch", help="Saves the result below to Prediction History."):
            log_history_entry({
                "asset_id": result["asset_id"], "timestamp": format_timestamp(result["timestamp"]),
                "failure_state": result["failure_state"], "confidence": result["confidence"],
                "remaining_useful_life": result["remaining_useful_life"],
            })
            st.session_state["single_prediction_saved"] = True
        if st.session_state.get("single_prediction_saved"):
            st.success("Saved to Prediction History.")

    stat_card_row([
        {"label": "Health status", "value": state.replace("_", " "), "tone": tone},
        {"label": "Confidence", "value": f"{result['confidence']:.1%}", "tone": tone},
        {"label": "Remaining useful life", "value": f"{result['remaining_useful_life']:.1f} h", "tone": tone},
        {"label": "Estimated replacement date", "value": estimated_replacement_date(result["remaining_useful_life"]), "tone": tone},
        {"label": "Prediction timestamp", "value": format_timestamp(result["timestamp"]), "tone": "accent"},
    ])

    if result.get("warnings"):
        for w in result["warnings"]:
            st.warning(w)

    st.divider()
    st.subheader("Condition assessment")
    diagnosis = result.get("fault_diagnosis")
    with st.container(border=True):
        st.markdown(
            _diagnosis_cell_html(state, diagnosis["predicted_fault"] if diagnosis else None),
            unsafe_allow_html=True,
        )
        if diagnosis:
            st.caption(
                "The classifier only stages health (Normal/Warning/Critical). This rule-based "
                "module (app/fault_diagnosis.py) runs on Warning and Critical predictions to "
                "estimate which historical failure mode the sensor pattern most resembles."
            )
            d1, d2 = st.columns(2)
            d1.markdown(
                f"**Diagnosis code**<br>{state_badge(diagnosis['predicted_fault'])}", unsafe_allow_html=True
            )
            d2.metric("Rule-engine confidence", f"{diagnosis['confidence']:.1%}")
            st.caption(f"Dominant channel: {diagnosis['dominant_channel']}")
            st.info(diagnosis["rationale"])
            st.dataframe(
                pd.DataFrame(
                    {"Channel": list(diagnosis["channel_deviation_scores"].keys()),
                     "Deviation (|z-score| from Normal baseline)": list(diagnosis["channel_deviation_scores"].values())}
                ),
                width="stretch", hide_index=True,
            )

    proba_df = pd.DataFrame(
        {"Class": list(result["class_probabilities"].keys()), "Probability": list(result["class_probabilities"].values())}
    )
    fig = px.bar(
        proba_df, x="Class", y="Probability", color="Class",
        color_discrete_map=CLASS_COLORS, title="Class probability distribution",
    )
    fig.update_layout(
        showlegend=False, yaxis_range=[0, 1], yaxis_title="Probability", xaxis_title="Predicted class",
        yaxis_tickformat=".0%",
    )
    st.plotly_chart(fig, width="stretch")

    if explanation:
        st.divider()
        st.subheader("Why the models predicted this")
        tab1, tab2 = st.tabs(["Classifier", "Regressor"])
        with tab1:
            st.caption(f"Top features pushing toward '{explanation['predicted_class_label']}':")
            st.dataframe(pd.DataFrame(explanation["classifier_top_features"]), width="stretch", hide_index=True)
            st.image(
                f"data:image/png;base64,{explanation['classifier_waterfall_png_base64']}",
                caption="SHAP waterfall — classifier",
            )
        with tab2:
            st.caption("Top features driving the predicted RUL:")
            st.dataframe(pd.DataFrame(explanation["regressor_top_features"]), width="stretch", hide_index=True)
            st.image(
                f"data:image/png;base64,{explanation['regressor_waterfall_png_base64']}",
                caption="SHAP waterfall — regressor",
            )


def build_batch_export_dataframe(result_df: pd.DataFrame) -> pd.DataFrame:
    df = add_priority_column(add_diagnosis_export_columns(result_df))
    if "RUL_Hours_Predicted" in df.columns:
        df["Estimated_Replacement_Date"] = df["RUL_Hours_Predicted"].apply(estimated_replacement_date)
    df["Model_Version"] = get_active_model_version()
    return df


def render_download_buttons(df: pd.DataFrame, xlsx_bytes: bytes, key_prefix: str) -> None:
    csv_col, xlsx_col = st.columns(2)
    csv_col.download_button(
        "Download Predictions (.csv)", data=df.to_csv(index=False),
        file_name="predictions.csv", mime="text/csv", key=f"{key_prefix}_csv", width="stretch",
        help="Plain data, all rows, no cell colors.",
    )
    xlsx_col.download_button(
        "Download Predictions (.xlsx)", data=xlsx_bytes,
        file_name="predictions.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key=f"{key_prefix}_xlsx", width="stretch",
        help="Same data, with red/yellow RUL highlighting preserved.",
    )


def read_uploaded_table(uploaded) -> pd.DataFrame:
    if uploaded.name.lower().endswith((".xlsx", ".xls")):
        return pd.read_excel(uploaded)
    return pd.read_csv(uploaded)


def render_retrain_section() -> None:
    st.subheader("Retrain Model")
    st.caption(
        "Upload a new labeled dataset (must include Failure_State and RUL_Hours columns) to retrain "
        "the health classifier and RUL regressor using the same preprocessing, feature engineering, "
        "feature selection, and training pipeline as the current production model."
    )

    if st.button("Close Retrain Panel", key="close_retrain_panel"):
        st.session_state["show_retrain_section"] = False
        st.rerun()

    retrain_uploaded = st.file_uploader(
        "Upload new training dataset (CSV or Excel)", type=["csv", "xlsx", "xls"], key="retrain_file_uploader",
        help="Needs the same columns as the original training set: Asset_ID, Machine_Model, Location, "
             "Vibration_mm_s, Temperature_C, Pressure_psi, Flow_Rate_m3_h, Failure_State, RUL_Hours.",
    )
    if retrain_uploaded is None:
        return

    try:
        with st.spinner("Reading uploaded dataset..."):
            retrain_df = read_uploaded_table(retrain_uploaded)
    except Exception as exc:
        st.error(f"Could not read the uploaded dataset: {exc}")
        return

    try:
        validation = retrain.validate_training_dataset(retrain_df)
    except Exception as exc:
        st.error(f"Could not validate the uploaded dataset: {exc}")
        return

    st.markdown("**1. Dataset preview**")
    with st.container(border=True):
        stat_card_row([
            {"label": "Rows", "value": f"{validation.n_rows:,}", "tone": "accent"},
            {"label": "Columns", "value": validation.n_cols, "tone": "accent"},
            {"label": "Missing values", "value": validation.n_missing, "tone": "warn" if validation.n_missing else "good"},
            {"label": "Duplicate rows", "value": validation.n_duplicates, "tone": "warn" if validation.n_duplicates else "good"},
        ])
        st.caption("First 10 rows:")
        st.dataframe(validation.preview, width="stretch", hide_index=True)

    st.markdown("**2. Validation**")
    with st.container(border=True):
        for err in validation.errors:
            st.error(err)
        for warn in validation.warnings:
            st.warning(warn)
        if validation.is_valid:
            st.success("Dataset passed validation. Ready to retrain.")
        else:
            return

    st.markdown("**3. Train**")
    if st.button("Start Retraining", type="primary", key="start_retraining_btn"):
        with st.status("Training in progress...", expanded=True) as status_box:
            stage_counter = {"n": 0}

            def _on_stage(name: str) -> None:
                stage_counter["n"] += 1
                status_box.update(label=f"Step {stage_counter['n']}/{len(retrain.TRAINING_STAGES)}: {name}")
                st.write(name)

            try:
                current_artifacts = load_artifacts()
                result = retrain.retrain_pipeline(
                    retrain_df, current_artifacts=current_artifacts,
                    models_root=config.MODELS_DIR, repo_dir=config.BASE_DIR,
                    dataset_version=retrain_uploaded.name, on_stage=_on_stage,
                )
            except Exception as exc:
                status_box.update(label="Retraining failed", state="error", expanded=True)
                st.error(f"Retraining failed: {exc}")
                return
            status_box.update(label=f"Retraining complete -- registered as {result.version}", state="complete", expanded=False)
        st.session_state["retrain_result"] = result
        st.session_state["retrain_decision"] = None

    result = st.session_state.get("retrain_result")
    if result is None:
        return
    for w in result.warnings:
        st.warning(w)

    training_date = result.artifacts.metadata.get("training_date", "unknown")
    stat_card_row([
        {"label": "Model version", "value": result.version, "tone": "accent"},
        {"label": "Training dataset size", "value": f"{result.n_training_rows:,} rows", "tone": "accent"},
        {"label": "Test set size", "value": f"{result.n_test_rows:,} rows", "tone": "accent"},
        {"label": "Last retrained", "value": format_timestamp(training_date) if training_date != "unknown" else "unknown", "tone": "accent"},
    ])

    st.markdown("**4. Model evaluation — current production vs. newly trained**")
    clf_comparison = pd.DataFrame({
        "Metric": ["Accuracy", "Precision", "Recall", "F1 Score"],
        "Current Production": [
            result.current_classifier_metrics["Accuracy"], result.current_classifier_metrics["Precision"],
            result.current_classifier_metrics["Recall"], result.current_classifier_metrics["F1_Score"],
        ],
        "Newly Trained": [
            result.new_classifier_metrics["Accuracy"], result.new_classifier_metrics["Precision"],
            result.new_classifier_metrics["Recall"], result.new_classifier_metrics["F1_Score"],
        ],
    })
    reg_comparison = pd.DataFrame({
        "Metric": ["RMSE", "MAE", "R2 Score"],
        "Current Production": [
            result.current_regressor_metrics["RMSE"], result.current_regressor_metrics["MAE"],
            result.current_regressor_metrics["R2_Score"],
        ],
        "Newly Trained": [
            result.new_regressor_metrics["RMSE"], result.new_regressor_metrics["MAE"],
            result.new_regressor_metrics["R2_Score"],
        ],
    })
    with st.container(border=True):
        e1, e2 = st.columns(2)
        with e1:
            st.caption("Classification metrics")
            st.dataframe(
                clf_comparison.style.format({"Current Production": "{:.4f}", "Newly Trained": "{:.4f}"}),
                width="stretch", hide_index=True,
            )
        with e2:
            st.caption("Regression metrics")
            st.dataframe(
                reg_comparison.style.format({"Current Production": "{:.4f}", "Newly Trained": "{:.4f}"}),
                width="stretch", hide_index=True,
            )

    st.markdown("**5. Decision**")
    is_better = retrain.recommend_replacement(result)
    decision = st.session_state.get("retrain_decision")

    if is_better:
        if decision is None:
            st.success(
                f"The newly trained model (version {result.version}) performs at least as well as the "
                "current production model on both classification and regression metrics, and better on "
                "at least one."
            )
            st.markdown("Replace the current production model with this one?")
            c1, c2 = st.columns(2)
            if c1.button(f"Yes, replace with {result.version}", type="primary", key="confirm_replace_btn", width="stretch"):
                try:
                    mr.promote_to_stable(config.MODELS_DIR, result.version)
                except Exception as exc:
                    st.error(f"Could not promote {result.version} to stable: {exc}")
                else:
                    load_artifacts.clear()
                    st.session_state["retrain_decision"] = "replaced"
                    st.rerun()
            if c2.button("No, keep current model", key="keep_current_btn", width="stretch"):
                st.session_state["retrain_decision"] = "kept"
                st.rerun()
        elif decision == "replaced":
            st.success(f"Production model replaced with version {result.version}. Future predictions will use it.")
        elif decision == "kept":
            st.info("Current production model retained. The new candidate model was saved but not promoted.")
    else:
        st.warning(
            f"The newly trained model (version {result.version}) does not clearly outperform the current "
            "production model, so it was saved as a candidate but not promoted."
        )

    st.markdown("**Training history**")
    history_df = retrain.training_history(config.MODELS_DIR)
    st.dataframe(history_df, width="stretch", hide_index=True)


def page_batch_prediction():
    page_header(
        "Batch Prediction",
        caption="Upload a CSV or Excel (.xlsx) file with columns: Asset_ID (optional), Timestamp (optional), "
                "Machine_Model, Location, Vibration_mm_s, Temperature_C, Pressure_psi, Flow_Rate_m3_h.",
    )

    uploaded = st.file_uploader(
        "Upload CSV or Excel", type=["csv", "xlsx", "xls"],
        help="Each row is one pump reading. Asset_ID and Timestamp are optional -- they're generated if missing.",
    )
    if uploaded is None:
        return

    try:
        with st.spinner("Reading uploaded file..."):
            input_df = read_uploaded_table(uploaded)
    except Exception as exc:
        st.error(f"Could not read the uploaded file: {exc}")
        return
    with st.container(border=True):
        st.caption(f"Preview -- {len(input_df):,} row(s) uploaded")
        st.dataframe(input_df, width="stretch", height=200, hide_index=True)

    run_col, retrain_col = st.columns(2)
    run_clicked = run_col.button("Run Batch Prediction", type="primary", width="stretch")
    retrain_clicked = retrain_col.button("Retrain Model", width="stretch")

    if retrain_clicked:
        st.session_state["show_retrain_section"] = not st.session_state.get("show_retrain_section", False)

    if run_clicked:
        try:
            with st.spinner(f"Scoring {len(input_df):,} row(s)..."):
                result_df = call_predict_csv(input_df.to_csv(index=False).encode("utf-8"), "batch.csv")
            if "Diagnosed_Fault" in result_df.columns:
                result_df["Diagnosed_Fault"] = result_df["Diagnosed_Fault"].fillna("")
            export_df = build_batch_export_dataframe(result_df)
            st.session_state["batch_result_df"] = result_df
            st.session_state["batch_result_xlsx"] = build_predictions_excel(export_df)
            st.session_state["batch_result_source_id"] = uploaded.file_id
            log_history_entry({
                "asset_id": f"batch ({len(result_df)} rows)", "timestamp": format_timestamp(pd.Timestamp.utcnow()),
                "failure_state": "-", "confidence": float("nan"), "remaining_useful_life": float("nan"),
            })
        except Exception as exc:
            st.error(f"Batch prediction failed: {exc}")
            return

    if st.session_state.get("show_retrain_section"):
        st.divider()
        render_retrain_section()

    result_df = st.session_state.get("batch_result_df")
    if result_df is None:
        return

    st.divider()
    if st.session_state.get("batch_result_source_id") != uploaded.file_id:
        st.warning(
            "These results are from a previously uploaded file, not the one shown in the preview "
            "above. Click **Run Batch Prediction** to score the currently uploaded file."
        )
    st.success(f"Scored {len(result_df):,} row(s).")

    severity = result_df["RUL_Hours_Predicted"].apply(rul_severity)
    critical_count = int((severity == "critical").sum())
    warning_count = int((severity == "warning").sum())
    total_count = len(result_df)

    stat_card_row([
        {"label": "🔴 Immediate replacement", "value": critical_count, "tone": "crit"},
        {"label": f"🟡 Replace within {WARNING_RUL_THRESHOLD_DAYS} days", "value": warning_count, "tone": "warn"},
        {"label": "📊 Total pumps analyzed", "value": total_count, "tone": "accent"},
    ])

    st.subheader("Results")
    st.caption(f"Rows with RUL ≤ {CRITICAL_RUL_THRESHOLD_HOURS}h are highlighted red and rows with "
               f"RUL ≤ {WARNING_RUL_THRESHOLD_HOURS}h are highlighted yellow below "
               "(CSV downloads can't carry cell colors -- pick Excel below to keep the highlighting). "
               "The **Priority** column repeats that signal as text so it doesn't rely on color alone. "
               "The **Diagnosis** column is populated for every row: Normal Operation for Normal, a "
               "likely developing issue for Warning, and a likely failure for Critical -- Warning and "
               "Critical both come from the existing fault diagnosis model output, never invented here.")
    display_df = add_priority_column(add_diagnosis_column(result_df))
    display_df, cap_notice = cap_for_onscreen_display(display_df)
    if cap_notice:
        st.caption(cap_notice)
    render_results_table_with_diagnosis(display_df, height_px=380)

    st.caption(
        "Downloads below include the same Priority and Diagnosis signal shown in the table above "
        "(as separate Likely Failure / Failure Description / Recommended Action columns), plus "
        "Estimated Replacement Date and Model Version."
    )
    export_df = build_batch_export_dataframe(result_df)
    render_download_buttons(export_df, st.session_state["batch_result_xlsx"], "whole")

    st.divider()
    st.subheader("Charts")
    col1, col2 = st.columns(2)
    with col1:
        counts = result_df["Failure_State_Predicted"].value_counts().reset_index()
        counts.columns = ["Failure_State", "Count"]
        fig = px.bar(counts, x="Failure_State", y="Count", color="Failure_State", color_discrete_map=CLASS_COLORS)
        fig.update_layout(
            showlegend=False, title="Predicted health status distribution",
            xaxis_title="Health status", yaxis_title="Pump count",
        )
        st.plotly_chart(fig, width="stretch")
    with col2:
        fig2 = px.histogram(result_df, x="RUL_Hours_Predicted", nbins=30, title="Predicted RUL distribution")
        fig2.update_layout(xaxis_title="Remaining useful life (hours)", yaxis_title="Pump count")
        fig2.update_traces(marker_color=ACCENT)
        st.plotly_chart(fig2, width="stretch")

    if "Confidence" in result_df.columns:
        fig_conf = px.histogram(result_df, x="Confidence", nbins=20, title="Prediction confidence distribution")
        fig_conf.update_layout(xaxis_title="Confidence", yaxis_title="Pump count", xaxis_range=[0, 1])
        fig_conf.update_traces(marker_color=ACCENT)
        st.plotly_chart(fig_conf, width="stretch")

    if {"Machine_Model", "Location", "Failure_State_Predicted"} <= set(result_df.columns):
        st.markdown("**Health status by machine model and location**")
        col3, col4 = st.columns(2)
        state_order = {"category_orders": {"Failure_State_Predicted": ["Normal", "Warning", "Critical"]}}
        with col3:
            machine_summary = (
                result_df.groupby(["Machine_Model", "Failure_State_Predicted"]).size().reset_index(name="Count")
            )
            fig_machine = px.bar(
                machine_summary, x="Machine_Model", y="Count", color="Failure_State_Predicted",
                color_discrete_map=CLASS_COLORS, barmode="group", **state_order,
            )
            fig_machine.update_layout(
                title="Machine-wise health summary", xaxis_title="Machine model", yaxis_title="Pump count",
                legend_title_text="Health status",
            )
            st.plotly_chart(fig_machine, width="stretch")
        with col4:
            location_summary = (
                result_df.groupby(["Location", "Failure_State_Predicted"]).size().reset_index(name="Count")
            )
            fig_location = px.bar(
                location_summary, x="Location", y="Count", color="Failure_State_Predicted",
                color_discrete_map=CLASS_COLORS, barmode="group", **state_order,
            )
            fig_location.update_layout(
                title="Location-wise health summary", xaxis_title="Location", yaxis_title="Pump count",
                legend_title_text="Health status",
            )
            st.plotly_chart(fig_location, width="stretch")

    diagnosed = result_df.get("Diagnosed_Fault")
    if diagnosed is not None and (diagnosed != "").any():
        st.subheader("Fault diagnosis for Warning & Critical rows")
        st.caption(
            "Only 4 of the 15 documented fault types can currently be diagnosed -- Bearing Failure, "
            "Motor Failure, Seal Failure, and Unknown Critical Fault -- because app/fault_diagnosis.py's "
            "rule engine only distinguishes those three sensor-deviation patterns today. This chart "
            "reflects that honestly rather than implying broader coverage; it will pick up any "
            "additional fault type automatically if the diagnosis module is ever extended."
        )
        fault_counts = diagnosed[diagnosed != ""].value_counts().reset_index()
        fault_counts.columns = ["Diagnosed_Fault", "Count"]
        fig3 = px.bar(
            fault_counts, x="Diagnosed_Fault", y="Count", color="Diagnosed_Fault",
            color_discrete_map=CLASS_COLORS,
        )
        fig3.update_layout(
            showlegend=False, title="Diagnosed fault distribution (Warning & Critical rows)",
            xaxis_title="Diagnosed fault", yaxis_title="Pump count",
        )
        st.plotly_chart(fig3, width="stretch")

    st.divider()
    st.subheader("Pumps scheduled for replacement")
    st.caption(f"Every row with RUL ≤ {WARNING_RUL_THRESHOLD_HOURS}h, sorted most urgent first.")
    replacement_df = build_replacement_schedule(result_df)
    if replacement_df.empty:
        st.info(f"No pumps currently require replacement within {WARNING_RUL_THRESHOLD_DAYS} days.")
    else:
        with st.container(border=True):
            st.dataframe(replacement_df, width="stretch", height=300, hide_index=True)


def page_prediction_history():
    page_header("Prediction History", caption="Every prediction this dashboard has made, two ways.")

    tab1, tab2 = st.tabs(["This session", "Audit trail (SQLite, all predictions)"])

    with tab1:
        history = st.session_state.get("history", [])
        if not history:
            st.info("No predictions made yet this session. Try Single Prediction or Batch Prediction.")
        else:
            st.dataframe(pd.DataFrame(history), width="stretch", hide_index=True)
            if st.button("Clear session history"):
                st.session_state["history"] = []
                st.rerun()

    with tab2:
        st.caption("Every prediction (single + batch), persisted to logs/audit_log.db and logs/audit_log.csv "
                   "-- includes model version and inference latency for each record.")
        col1, col2 = st.columns([3, 1])
        asset_filter = col1.text_input("Filter by Asset ID (optional)", value="", help="Leave blank to see the most recent predictions across all assets.")
        limit = col2.number_input("Limit", min_value=10, max_value=1000, value=100, step=10)
        try:
            with st.spinner("Loading audit trail..."):
                if asset_filter.strip():
                    resp = requests.get(f"{API_URL}/audit/asset/{asset_filter.strip()}", params={"limit": limit}, timeout=10)
                else:
                    resp = requests.get(f"{API_URL}/audit/recent", params={"limit": limit}, timeout=10)
                resp.raise_for_status()
                records = resp.json()["records"]
        except Exception as exc:
            st.error(f"Could not load audit trail: {exc}")
            return

        if not records:
            st.info("No audit records yet. Make a prediction via the API or dashboard first.")
            return

        audit_df = pd.DataFrame(records)
        parsed_features = audit_df.get("input_features", pd.Series([""] * len(audit_df))).apply(
            lambda v: json.loads(v) if isinstance(v, str) and v else (v if isinstance(v, dict) else {})
        )
        audit_df["Machine_Model"] = parsed_features.apply(lambda d: d.get("Machine_Model", ""))
        audit_df["Location"] = parsed_features.apply(lambda d: d.get("Location", ""))
        audit_df["_raw_timestamp"] = pd.to_datetime(audit_df["timestamp"], errors="coerce", utc=True)
        if "rul_hours" in audit_df.columns:
            audit_df["Estimated_Replacement_Date"] = audit_df["rul_hours"].apply(estimated_replacement_date)

        st.caption(f"Filters below apply within the {len(audit_df):,} record(s) fetched above.")
        f1, f2, f3, f4 = st.columns(4)
        machine_options = ["All"] + sorted({m for m in audit_df["Machine_Model"] if m})
        location_options = ["All"] + sorted({loc for loc in audit_df["Location"] if loc})
        machine_choice = f1.selectbox("Machine Model", machine_options)
        location_choice = f2.selectbox("Location", location_options)
        status_choice = f3.selectbox("Health Status", ["All", "Normal", "Warning", "Critical"])

        valid_dates = audit_df["_raw_timestamp"].dropna()
        if not valid_dates.empty:
            min_date, max_date = valid_dates.min().date(), valid_dates.max().date()
            date_range = f4.date_input("Date range", value=(min_date, max_date), min_value=min_date, max_value=max_date)
        else:
            date_range = f4.date_input("Date range", value=(), disabled=True, help="No parseable timestamps in this window.")

        filtered = audit_df
        if machine_choice != "All":
            filtered = filtered[filtered["Machine_Model"] == machine_choice]
        if location_choice != "All":
            filtered = filtered[filtered["Location"] == location_choice]
        if status_choice != "All" and "predicted_state" in filtered.columns:
            filtered = filtered[filtered["predicted_state"] == status_choice]
        if isinstance(date_range, (tuple, list)) and len(date_range) == 2:
            start_date, end_date = date_range
            in_range = filtered["_raw_timestamp"].dt.date.between(start_date, end_date)
            filtered = filtered[in_range | filtered["_raw_timestamp"].isna()]

        filtered = filtered.drop(columns=["_raw_timestamp"], errors="ignore")
        if "timestamp" in filtered.columns:
            filtered = filtered.assign(timestamp=filtered["timestamp"].apply(format_timestamp))

        st.caption(f"{len(filtered):,} of {len(audit_df):,} fetched record(s) match the filters above.")
        st.dataframe(filtered, width="stretch", height=400, hide_index=True)
        st.download_button(
            "Download as CSV", data=filtered.to_csv(index=False),
            file_name="audit_trail_export.csv", mime="text/csv", key="audit_trail_csv",
        )


def page_about():
    page_header("About")
    st.markdown(
        f"""
        **{config.APP_NAME}** v{config.APP_VERSION}

        End-to-end predictive maintenance system for industrial centrifugal
        pumps: EDA → feature engineering → model training/comparison
        (Decision Tree, Random Forest, XGBoost, LightGBM) → persistence →
        this FastAPI + Streamlit serving layer.

        **Stack:** pandas, scikit-learn, XGBoost, LightGBM, FastAPI, Streamlit, Docker.

        **Honest caveat:** the classifier's Normal/Warning/Critical predictions
        are reliable (non-overlapping sensor bands in training data). Component-level
        fault typing (Bearing/Motor/Seal_Failure) is **not** a classifier output
        anymore — those ~20 historical examples, from a single pump, have
        near-identical sensor signatures across all three failure modes (all four
        channels pegged near their extreme simultaneously), too few and too
        undifferentiated to fit or validate a reliable model on. Instead,
        `app/fault_diagnosis.py` runs a transparent rule engine (on Warning and
        Critical predictions) using the standard reliability-engineering heuristic —
        vibration → bearing, temperature → motor, pressure/flow drop → seal —
        and honestly reports "Unknown_Critical_Fault" when no channel clearly
        dominates rather than guessing. The dashboard frames the same underlying
        diagnosis as a developing issue to monitor at Warning severity, and a
        likely failure at Critical severity. Recalibrate or replace it with a trained
        model once real, physically-differentiated failure data is collected
        across multiple assets.

        **API backend:** `{API_URL}` (override with the `PDM_API_URL` environment variable).
        """
    )


PAGES = {
    "Home": page_home,
    "Project Overview": page_overview,
    "Single Prediction": page_single_prediction,
    "Batch Prediction": page_batch_prediction,
    "Prediction History": page_prediction_history,
    "About": page_about,
}

inject_custom_css()

st.sidebar.markdown("## PdM Dashboard")

_health = api_health()
_status_class = "online" if _health else "offline"
_status_text = "API online" if _health else "API offline"
st.sidebar.markdown(
    f'<div class="pdm-sidebar-status"><span class="pdm-dot {_status_class}"></span>{_status_text}</div>',
    unsafe_allow_html=True,
)

selection = st.sidebar.radio("Navigate", list(PAGES.keys()), label_visibility="collapsed")
st.sidebar.divider()
st.sidebar.caption(f"API: {API_URL}")

PAGES[selection]()
